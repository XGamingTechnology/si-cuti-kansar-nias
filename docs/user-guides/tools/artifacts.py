"""Style genuine DOCX and verify export structure/render bounds. No application access."""
from pathlib import Path
import hashlib
import json
import sys
import zipfile
import fitz
from docx import Document
from docx.shared import Inches, Pt, RGBColor


def style(path):
    d = Document(path)
    for sec in d.sections:
        sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
        sec.top_margin = sec.bottom_margin = Inches(.55)
        sec.left_margin = sec.right_margin = Inches(.65)
    normal = d.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(11)
    for name, size in [('Title', 26), ('Heading 1', 23), ('Heading 2', 18), ('Heading 3', 14)]:
        target = next((s for s in d.styles if s.name.lower() == name.lower()), None)
        if target is None:
            continue
        target.font.color.rgb = RGBColor.from_string('103B55')
        target.font.size = Pt(size)
        target.paragraph_format.keep_with_next = True
    for shape in d.inline_shapes:
        factor = min(Inches(10.2) / shape.width, Inches(6.25) / shape.height)
        shape.width, shape.height = int(shape.width * factor), int(shape.height * factor)
    for p in d.paragraphs:
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.keep_with_next = True
        if p.text.startswith('Gambar '):
            p.paragraph_format.keep_together = True
            for run in p.runs:
                run.font.size = Pt(9)
    for t in d.tables:
        t.autofit = True
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
    for sec in d.sections:
        p = sec.footer.paragraphs[0]
        p.text = 'SI CUTI · Panduan 1.0 · 18 Oktober 2026 · Data contoh lokal'
        p.runs[0].font.size = Pt(8)
    d.save(path)


def validate(output):
    results = []
    for role in ['pegawai', 'admin']:
        pdf_path = output / f'panduan-{role}.pdf'
        pdf = fitz.open(pdf_path)
        images = 0
        for n, page in enumerate(pdf, 1):
            assert page.get_text().strip() or page.get_images(), f'Blank page {n}'
            for b in page.get_text('blocks'):
                rect = fitz.Rect(b[:4])
                assert rect.x0 >= -1 and rect.y0 >= -1 and rect.x1 <= page.rect.width+1 and rect.y1 <= page.rect.height+1, f'Text overflow {role}:{n}'
            for image in page.get_images():
                for rect in page.get_image_rects(image[0]):
                    assert rect.x0 >= 0 and rect.y0 >= 0 and rect.x1 <= page.rect.width+1 and rect.y1 <= page.rect.height+1, f'Image overflow {role}:{n}'
            images += len(page.get_images())
        docx = output / f'panduan-{role}.docx'
        d = Document(docx)
        with zipfile.ZipFile(docx) as z:
            assert z.testzip() is None
            assert '[Content_Types].xml' in z.namelist()
        assert images == len(d.inline_shapes), f'Image mismatch {role}'
        for shape in d.inline_shapes:
            assert shape.width <= Inches(10.21) and shape.height <= Inches(6.26)
        assert any('4c341b0b67bbeed0ea366fb22b3d369c47677b39' in p.text for p in d.paragraphs)
        results.append({'role': role, 'pdfPages': len(pdf), 'pdfFigures': images, 'docxInlineFigures': len(d.inline_shapes), 'pdfBounds': 'PASS', 'docxZipAndExtents': 'PASS', 'docxNativeWordRender': 'NOT_AVAILABLE', 'docxContentRoundtripPreview': 'Generated; not a native layout proof'})
    (output/'artifact-validation.json').write_text(json.dumps(results, indent=2)+'\n')
    (output/'SHA256SUMS.txt').write_text(''.join(f'{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.name}\n' for p in sorted(output.iterdir()) if p.suffix in ['.pdf', '.docx']))
    print(json.dumps(results, indent=2))


if __name__ == '__main__':
    if sys.argv[1] == 'style':
        style(Path(sys.argv[2]))
    else:
        validate(Path(sys.argv[2]))
